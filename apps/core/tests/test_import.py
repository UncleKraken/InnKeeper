from decimal import Decimal

from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from apps.core import importer
from apps.core.testing import PASSWORD, HotelTestCase
from apps.frontdesk.models import Guest
from apps.inventory.models import StockItem
from apps.outlets.models import Item, Outlet


class ImportTests(HotelTestCase):
    def test_menu_with_albanian_headers_semicolons_and_comma_decimals(self):
        raw = "Pika;Kategoria;Emri;Çmimi;Emri anglisht\nRestaurant;Mains;Tavë kosi;12,50;Baked lamb\nBar;Kafe;Espresso;1,5;\n"
        parsed = importer.parse("menu", raw.encode("utf-8-sig"))
        self.assertEqual(parsed.errors, [])
        stats = importer.apply(parsed, self.manager)
        self.assertEqual(stats, {"created": 1, "updated": 1})  # Tavë kosi existed
        self.dish.refresh_from_db()
        self.assertEqual((self.dish.price, self.dish.name_en), (Decimal("12.50"), "Baked lamb"))
        self.assertTrue(
            Item.objects.filter(name="Espresso", category__outlet__name="Bar", price=Decimal("1.50")).exists()
        )
        self.assertEqual(Outlet.objects.filter(name="Bar").count(), 1)

    def test_bad_rows_are_reported_and_skipped(self):
        raw = b"outlet,category,name,price\nRestaurant,Mains,Soup,abc\nRestaurant,Mains,Salad,5\n"
        parsed = importer.parse("menu", raw)
        self.assertEqual(len(parsed.rows), 1)
        self.assertIn("2", parsed.errors[0])

    def test_missing_columns(self):
        parsed = importer.parse("menu", b"name,price\nSoup,4\n")
        self.assertTrue(parsed.errors)
        self.assertEqual(parsed.rows, [])

    def test_stock_sets_quantity_with_a_movement(self):
        raw = "name;group;unit;in stock;cost;supplier\nBeer 0.5L;Drinks;bottle;48;0,85;Birra Korça\n"
        importer.apply(importer.parse("stock", raw.encode()), self.manager)
        item = StockItem.objects.get(name="Beer 0.5L")
        self.assertEqual(
            (item.on_hand, item.unit, item.group, item.supplier.name),
            (Decimal("48"), "bottle", "drinks", "Birra Korça"),
        )
        self.assertEqual(item.moves.count(), 1)

    def test_guests_with_dates_and_cp1252(self):
        raw = "first_name;last_name;nationality;document_type;document_number;date_of_birth\nÉlise;Dupont;FR;Passport;12AB;31.12.1990\n"
        importer.apply(importer.parse("guests", raw.encode("cp1252")), self.manager)
        g = Guest.objects.get(last_name="Dupont")
        self.assertEqual((g.first_name, g.document_type, str(g.date_of_birth)), ("Élise", "passport", "1990-12-31"))

    def test_upload_preview_and_confirm(self):
        self.client.login(username="boss", password=PASSWORD)
        self.assertEqual(self.client.get(reverse("core:settings") + "?tab=import").status_code, 200)
        self.assertEqual(self.client.get(reverse("core:import_template", args=["menu"])).status_code, 200)
        f = SimpleUploadedFile("menu.csv", b"outlet;category;name;price\nRestaurant;Mains;Fergese;9\n", "text/csv")
        page = self.client.post(reverse("core:import_file"), {"kind": "menu", "file": f})
        self.assertContains(page, "Fergese")
        self.assertFalse(Item.objects.filter(name="Fergese").exists())
        self.client.post(reverse("core:import_file"), {"kind": "menu", "file": page.context["file"], "confirm": "1"})
        self.assertTrue(Item.objects.filter(name="Fergese").exists())

    def test_broken_xlsx_is_explained(self):
        self.client.login(username="boss", password=PASSWORD)
        f = SimpleUploadedFile("menu.xlsx", b"PK\x03\x04rest", "application/octet-stream")
        r = self.client.post(reverse("core:import_file"), {"kind": "menu", "file": f}, follow=True)
        self.assertEqual(r.context["parsed"].rows, [])
        self.assertIn(".xlsx", r.context["parsed"].errors[0])

    def test_old_xls_is_explained(self):
        self.client.login(username="boss", password=PASSWORD)
        f = SimpleUploadedFile("menu.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest", "application/vnd.ms-excel")
        r = self.client.post(reverse("core:import_file"), {"kind": "menu", "file": f}, follow=True)
        self.assertContains(r, ".xlsx")


def _xlsx(rows) -> bytes:
    import io

    from openpyxl import Workbook

    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


class ExcelImportTests(HotelTestCase):
    def test_menu_from_a_real_workbook_with_numbers_and_blank_lines(self):
        raw = _xlsx([[None], ["Pika", "Kategoria", "Emri", "Çmimi"], ["Bar", "Kafe", "Espresso", 1.5], [None, None], ["Bar", "Kafe", "Makiato", 2]])
        parsed = importer.parse("menu", raw)
        self.assertEqual(parsed.errors, [])
        self.assertEqual([(r["name"], r["price"]) for r in parsed.rows], [("Espresso", "1.50"), ("Makiato", "2.00")])

    def test_guest_dates_as_real_excel_dates(self):
        from datetime import datetime

        raw = _xlsx([["first_name", "last_name", "date_of_birth"], ["Ana", "Leka", datetime(1991, 3, 4)]])
        importer.apply(importer.parse("guests", raw), self.manager)
        self.assertEqual(str(Guest.objects.get(last_name="Leka").date_of_birth), "1991-03-04")

    def test_upload_preview_and_confirm_xlsx(self):
        self.client.login(username="boss", password=PASSWORD)
        raw = _xlsx([["outlet", "category", "name", "price"], ["Restaurant", "Mains", "Fergese", 9]])
        f = SimpleUploadedFile("menu.xlsx", raw, "application/octet-stream")
        page = self.client.post(reverse("core:import_file"), {"kind": "menu", "file": f})
        self.assertContains(page, "Fergese")
        self.assertTrue(page.context["file"].endswith(".xlsx"))
        self.client.post(reverse("core:import_file"), {"kind": "menu", "file": page.context["file"], "confirm": "1"})
        self.assertTrue(Item.objects.filter(name="Fergese", price=Decimal("9")).exists())

    def test_template_is_a_workbook_that_imports_back(self):
        self.client.login(username="boss", password=PASSWORD)
        for kind in importer.KINDS:
            r = self.client.get(reverse("core:import_template", args=[kind]))
            self.assertEqual(r["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            parsed = importer.parse(kind, r.content)
            self.assertEqual((parsed.errors, parsed.rows, parsed.unknown), ([], [], []), kind)
        r = self.client.get(reverse("core:import_template", args=["menu"]) + "?format=csv")
        self.assertIn(b"outlet;category", r.content)


class ExcelExportTests(HotelTestCase):
    def _book(self, content):
        import io

        from openpyxl import load_workbook

        return load_workbook(io.BytesIO(content))

    def test_ledger_xlsx_has_numbers_and_no_formulas(self):
        from apps.finance.models import Folio
        from apps.frontdesk import services

        folio, _ = Folio.objects.get_or_create(reservation=self.make_reservation())
        services.add_charge(folio, description='=HYPERLINK("http://x")', quantity=1, unit_price=Decimal("12.50"), user=self.manager)
        self.client.login(username="boss", password=PASSWORD)
        r = self.client.get(reverse("finance:ledger") + "?format=xlsx")
        self.assertEqual(r.status_code, 200)
        ws = self._book(r.content).active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        row = next(x for x in rows if x[3] and "HYPERLINK" in x[3])
        self.assertEqual(ws.cell(row=rows.index(row) + 2, column=4).data_type, "s")
        self.assertEqual(row[6], 12.5)

    def test_sales_stock_and_register_download(self):
        self.client.login(username="boss", password=PASSWORD)
        for url in (reverse("finance:sales"), reverse("inventory:stock"), reverse("frontdesk:guest_register")):
            r = self.client.get(url + "?format=xlsx")
            self.assertEqual(r.status_code, 200, url)
            self.assertTrue(r.content.startswith(b"PK"), url)
        self.assertEqual(self._book(self.client.get(reverse("finance:sales") + "?format=xlsx").content).sheetnames.__len__(), 4)
